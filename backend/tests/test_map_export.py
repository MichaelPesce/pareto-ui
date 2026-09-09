import copy
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from app.internal.ExcelApi import WriteMapDataToExcel, determineConnectionsFromArcs


class MapExportTests(unittest.TestCase):
    def test_expansion_defaults_survive_reorder_and_only_fill_new_entries(self):
        data = {'StorageSites': {'S1': {}, 'S2': {}}, 'SWDSites': {'K1': {}, 'K2': {}},
                'TreatmentSites': {'R1': {}}, 'NetworkNodes': {'N1': {}},
                'connections': {'all_connections': {'N1': ['K1']}}}
        with tempfile.TemporaryDirectory() as tmp:
            target = str(Path(tmp) / 'map')
            WriteMapDataToExcel(data, target)
            wb = load_workbook(target + '.xlsx')
            self.assertEqual(wb['StorageExpansionCost']['C3'].value, 2)
            self.assertEqual(wb['DisposalExpansionCost']['C3'].value, 1000)
            self.assertEqual(wb['DisposalExpansionLeadTime']['C3'].value, 45)
            self.assertIsNone(wb['DisposalExpansionLeadTime']['D3'].value)
            self.assertEqual(wb['PipelineExpansionLeadTime_Capac']['D3'].value, 1)
            self.assertIsNone(wb['PipelineCapexCapacityBased']['D3'].value)
            wb['StorageExpansionCost']['C3'] = 17
            wb['DisposalExpansionCost']['C3'] = 12345
            wb['DisposalExpansionLeadTime']['C3'] = 0
            wb['TreatmentExpansionCost']['D3'] = 123
            wb['PipelineCapexCapacityBased']['D3'] = 456
            wb.save(target + '.xlsx'); wb.close()
            changed = copy.deepcopy(data)
            changed['StorageSites'] = {'S2': {}, 'S1': {}, 'S3': {}}
            changed['SWDSites'] = {'K2': {}, 'K1': {}}
            for _ in range(2):
                WriteMapDataToExcel(changed, target, target + '.xlsx')
                wb = load_workbook(target + '.xlsx')
                self.assertEqual(wb['StorageExpansionCost']['C3'].value, 2)
                self.assertEqual(wb['StorageExpansionCost']['C4'].value, 17)
                self.assertEqual(wb['StorageExpansionCost']['C5'].value, 2)
                self.assertEqual(wb['DisposalExpansionCost']['C4'].value, 12345)
                self.assertEqual(wb['DisposalExpansionLeadTime']['C4'].value, 0)
                self.assertEqual(wb['TreatmentExpansionCost']['D3'].value, 123)
                self.assertEqual(wb['PipelineCapexCapacityBased']['D3'].value, 456)
                wb.close()
            changed['StorageSites'] = {}
            WriteMapDataToExcel(changed, target, target + '.xlsx')
            wb = load_workbook(target + '.xlsx')
            self.assertTrue(all(value is None for row in list(wb['StorageExpansionCost'].values)[1:] for value in row))
            wb.close()

    def test_index_only_optional_arc_sheet_is_empty_model_data(self):
        from app.internal.get_data import get_data
        with tempfile.TemporaryDirectory() as tmp:
            target = str(Path(tmp) / 'map')
            WriteMapDataToExcel({'NetworkNodes': {'N1': {}}, 'connections': {'all_connections': {}}}, target)
            wb = load_workbook(target + '.xlsx')
            wb['NOA']['A2'] = 'NetworkNodes'
            wb['NOA']['A3'] = 'N1'
            wb['ReuseMinimum']['A2'] = 'ReuseOptions'
            wb['ReuseMinimum']['B2'] = 'T01'
            wb['BeneficialReuseCost']['A2'] = 'ReuseOptions'
            wb['BeneficialReuseCost']['B2'] = 'VALUE'
            wb.save(target + '.xlsx'); wb.close()
            _, parameters, _ = get_data(target + '.xlsx')
            self.assertEqual(parameters['NOA'], {})
            self.assertEqual(parameters['ReuseMinimum'], {})
            self.assertEqual(parameters['BeneficialReuseCost'], {})

    def test_reverse_and_bidirectional_flow_receive_segment_metadata(self):
        data = {'arcs': {'pipe': {'diameter': 'D4', 'lengths': [0, 5], 'nodes': [
            {'name': 'A', 'outgoing_nodes': ['B']},
            {'name': 'B', 'outgoing_nodes': ['A']},
            {'name': 'C', 'outgoing_nodes': ['B']},
        ]}}}
        actual = determineConnectionsFromArcs(data, {'D4': 4}, {'D4': 14286})['connections']
        self.assertEqual(set(actual['connection_metadata']), {'A::B', 'B::A', 'C::B'})
        self.assertEqual(actual['connection_metadata']['C::B'], {
            'pipeline_length': 5, 'pipeline_diameter': 4, 'pipeline_capacity': 14286})
        self.assertEqual(actual['connection_metadata']['A::B']['pipeline_length'], 0)


if __name__ == '__main__':
    unittest.main()
